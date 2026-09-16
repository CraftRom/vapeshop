from __future__ import annotations
import io, re, secrets, string
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook, Workbook

SKU_RE = re.compile(r'^[A-Z0-9][A-Z0-9._-]{2,31}$')

def normalize_sku(value) -> str | None:
    if value is None: return None
    sku=str(value).strip().upper()
    return sku if SKU_RE.fullmatch(sku) else None

def category_code(name: str) -> str:
    chars=''.join(c for c in str(name).upper() if c in string.ascii_uppercase+string.digits)
    return (chars[:4] or 'GEN').ljust(3,'X')

async def generate_sku(repo, category_name='GEN') -> str:
    prefix=f'ELF-{category_code(category_name)}-'
    alphabet=string.ascii_uppercase+string.digits
    for _ in range(50):
        sku=prefix+''.join(secrets.choice(alphabet) for _ in range(8))
        if not await repo.get_product_by_sku(sku): return sku
    raise RuntimeError('Не вдалося згенерувати унікальний артикул')

def dec(v):
    try: return Decimal(str(v)) if v not in (None,'') else None
    except (InvalidOperation, ValueError): return None

def parse_salesdrive_xlsx(raw: bytes):
    wb=load_workbook(io.BytesIO(raw), read_only=True, data_only=True); ws=wb.active
    rows=ws.iter_rows(values_only=True); headers=[str(x or '').strip() for x in next(rows)]
    idx={h:i for i,h in enumerate(headers)}
    required={'Товар/Послуга','Категорія'}
    if not required.issubset(idx): raise ValueError('Файл не схожий на експорт SalesDrive: немає колонок Товар/Послуга або Категорія')
    out=[]
    for r in rows:
        name=str(r[idx['Товар/Послуга']] or '').strip()
        if not name: continue
        def get(k): return r[idx[k]] if k in idx and idx[k] < len(r) else None
        out.append({'name':name,'sku':get('SKU'),'category':str(get('Категорія') or 'Без категорії').strip(),
                    'description':get('Опис'),'photo_url':get('Зображення'),'regular_price':dec(get('Ціна')),
                    'discount_price':dec(get('Ціна зі знижкою')),'discount':dec(get('Знижка'))})
    return out

def export_xlsx(products):
    wb=Workbook(); ws=wb.active; ws.title='Products'
    ws.append(['SKU','Товар','Категорія','Ціна','Стара ціна','Залишок','Опис','Зображення','Активний'])
    for p in products: ws.append([p.sku,p.name,p.category_name,float(p.price),float(p.old_price) if p.old_price is not None else None,p.stock,p.description,p.photo_url,'Так' if p.is_active else 'Ні'])
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(48,max(12,max(len(str(c.value or '')) for c in col)+2))
    b=io.BytesIO(); wb.save(b); return b.getvalue()
