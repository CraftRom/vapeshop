from io import BytesIO
from openpyxl import Workbook, load_workbook
from shop.services.product_io import normalize_sku, parse_salesdrive_xlsx, export_xlsx
from shop.entities import Product
from decimal import Decimal

def check(v,m):
    if not v: raise AssertionError(m)
check(normalize_sku('abc-123')=='ABC-123','normalize SKU')
check(normalize_sku('a b') is None,'reject short/invalid SKU')
wb=Workbook(); ws=wb.active; ws.append(['Товар/Послуга','SKU','Ціна','Знижка','Ціна зі знижкою','Опис','Зображення','Категорія']); ws.append(['Test',None,100,10,90,'Desc',None,'Cat']); b=BytesIO(); wb.save(b)
r=parse_salesdrive_xlsx(b.getvalue()); check(len(r)==1 and r[0]['discount_price']==Decimal('90'),'parse SalesDrive XLSX')
p=Product(id=1,category_id=1,name='Test',price=Decimal('90'),sku='ELF-CAT-12345678',old_price=Decimal('100'),category_name='Cat')
out=export_xlsx([p]); w=load_workbook(BytesIO(out),read_only=True,data_only=True); check(w.active.cell(2,1).value=='ELF-CAT-12345678','export SKU')
print('product IO: OK')
